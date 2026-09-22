# -*- coding: utf-8 -*-
"""Excel 导出：最终分配表 / 共用清单 / 缺口参考 / 各班归置表 / 试穿记录。"""
import io
import base64

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from . import schemas, conflict_checker

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.active
    ws.title = title
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
    if widths:
        from openpyxl.utils import get_column_letter
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


def _wb_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def final_allocation_xlsx(store, members, equipment):
    """最终分配表：每名队员礼服/马靴、共用人、位置、备注、来源、锁定。"""
    member_by_id = {m["member_id"]: m for m in members}
    eq_by_id = {e["equipment_id"]: e for e in equipment}

    # 共用人
    from collections import defaultdict
    u_share, b_share = defaultdict(list), defaultdict(list)
    for a in store["allocations"]:
        if a.get("uniform_id"):
            u_share[a["uniform_id"]].append(member_by_id.get(a["member_id"], {}).get("name", ""))
        if a.get("boot_id"):
            b_share[a["boot_id"]].append(member_by_id.get(a["member_id"], {}).get("name", ""))

    wb = Workbook()
    rows = []
    for a in store["allocations"]:
        m = member_by_id.get(a["member_id"], {})
        uid = a.get("uniform_id")
        bid = a.get("boot_id")
        ue = eq_by_id.get(uid, {}) if uid else {}
        be = eq_by_id.get(bid, {}) if bid else {}
        rows.append([
            m.get("name", ""), m.get("duty_class", ""),
            f"{uid} / {ue.get('size_code', '')}" if uid else "",
            "、".join(n for n in u_share.get(uid, []) if n != m.get("name")) if uid else "",
            f"{bid} / {be.get('size_code', '')}" if bid else "",
            "、".join(n for n in b_share.get(bid, []) if n != m.get("name")) if bid else "",
            ue.get("location", ""),
            "锁定" if a.get("locked") else a.get("source", ""),
            a.get("fit_explanation", ""),
        ])
    _sheet(wb, "最终分配", ["姓名", "班级", "礼服ID/尺码", "礼服共用人", "马靴ID/尺码", "马靴共用人",
                           "礼服位置", "状态", "系统解释"], rows,
           [10, 10, 20, 16, 16, 16, 12, 12, 30])
    return _wb_bytes(wb)


def shared_list_xlsx(store, members, cfg):
    """共用清单：按实体装备汇总两个使用人、班级和风险等级。"""
    class_matrix = cfg.get("class_matrix", {})
    n2m = {m["name"]: m for m in members}

    wb = Workbook()
    rows = []
    for s in store["last_result"].get("shared", []) if store["last_result"] else []:
        names = s["members"]
        cs = [(n2m.get(n, {}).get("duty_class") or "") for n in names]
        level = _risk_level(cs, class_matrix)
        rows.append([s["equipment_id"], "礼服" if s["category"] == "uniform" else "马靴",
                     "、".join(names), "、".join(cs), level])
    _sheet(wb, "共用清单", ["装备ID", "类别", "使用人", "班级", "风险等级"], rows,
           [16, 8, 22, 18, 10])
    return _wb_bytes(wb)


def _risk_level(classes, class_matrix):
    if len(classes) < 2:
        return "低"
    if classes[0] == classes[1]:
        return "禁止(同班)"
    lvl = class_matrix.get((classes[0], classes[1])) or class_matrix.get((classes[1], classes[0]))
    return {"high": "高", "mid": "中", "low": "低"}.get(lvl, "中")


def shortage_xlsx(store):
    """缺口与待采购参考。"""
    wb = Workbook()
    rows = []
    for s in (store["last_result"].get("shortage", []) if store["last_result"] else []):
        rows.append(["礼服" if s["category"] == "uniform" else "马靴", s["size_code"],
                     s["available"], s["max_cover"], s["demand"],
                     "、".join(s["affected"]), s["gap"], s["add_items"],
                     "；".join(f'{x["name"]}→{x["suggest"]}' for x in s["suggestions"])])
    _sheet(wb, "缺口参考", ["类别", "尺码", "可用实体", "最多覆盖", "仅需该码人数",
                          "受影响人", "缺口人数", "建议新增件数", "建议试穿"], rows,
           [8, 10, 10, 10, 12, 26, 10, 12, 34])
    return _wb_bytes(wb)


def placement_xlsx(store, equipment):
    """各班装备归置表：柜内礼服和马靴摆放顺序。"""
    wb = Workbook()
    rows = []
    uniforms = [e for e in equipment if e["category"] == "uniform"]
    uniforms.sort(key=lambda e: (e["gender"], schemas.parse_uniform_size(e["size_code"])))
    boots = [e for e in equipment if e["category"] == "boots"]
    boots.sort(key=lambda e: int(e["size_code"]) if str(e["size_code"]).isdigit() else 0)

    for e in uniforms:
        rows.append(["礼服", e["gender"], e["size_code"], e["equipment_id"], e["location"],
                     e.get("paired_belt_id", "")])
    for e in boots:
        rows.append(["马靴", "", e["size_code"], e["equipment_id"], e["location"], ""])
    _sheet(wb, "归置表", ["类别", "性别", "尺码", "装备ID", "位置", "配套腰带"], rows,
           [8, 8, 10, 16, 14, 12])
    return _wb_bytes(wb)


def fit_records_xlsx(store, members):
    """试穿记录备份。"""
    member_by_id = {m["member_id"]: m for m in members}
    wb = Workbook()
    rows = []
    for f in store["fit_records"]:
        rows.append([member_by_id.get(f["member_id"], {}).get("name", ""),
                     "礼服" if f["category"] == "uniform" else "马靴",
                     f["size_code"], f["fit_level"], f["fit_score"],
                     f["detail_note"], f["recorded_at"]])
    _sheet(wb, "试穿记录", ["姓名", "类别", "尺码", "试穿等级", "分值", "备注", "时间"], rows,
           [10, 8, 10, 10, 8, 30, 20])
    return _wb_bytes(wb)


def export_all(store, members, equipment, cfg):
    """返回 {文件名: bytes} 的导出包。"""
    return {
        "最终分配表.xlsx": final_allocation_xlsx(store, members, equipment),
        "共用清单.xlsx": shared_list_xlsx(store, members, cfg),
        "缺口与待采购参考.xlsx": shortage_xlsx(store),
        "各班装备归置表.xlsx": placement_xlsx(store, equipment),
        "试穿记录.xlsx": fit_records_xlsx(store, members),
    }
